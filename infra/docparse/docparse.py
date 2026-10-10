"""RUVO's document sandbox: scans a PDF, Excel workbook, CSV or JSON document that RUVO found on the web.

Documents are hostile input (active content, parser exploits, decompression bombs), so RUVO's API and worker never
parse them: each one is read here, in a container per document with no network, a read-only filesystem, and memory,
CPU and process limits.

Run: docparse < document    the document on stdin, one JSON object on stdout

The result always has these keys:
  verdict  "clean" (may be parsed), "risky" (active content: metadata only, never parsed) or "rejected" (dropped)
  reason   why the document was rejected, in plain words; else null
  type     what the bytes are, whatever the name or declared type says: "pdf", "xlsx", "csv", "json", "html" (a web
           page, always rejected) or null (anything else)
  flags    what the scan found: active content ("javascript", "open_action", "auto_action", "launch", "embedded_file",
           "rich_media", "xfa", "submit_form", "remote_goto" in a PDF; "macros", "external_link", "ole_object", "dde"
           in a workbook) and "encrypted"
  meta     what the file says about itself, {} when rejected. A PDF: title, author, subject, keywords, created,
           modified (ISO 8601) and pages. A workbook: the same but pages, plus sheets [{name, hidden, rows, columns}].
           A CSV: encoding, delimiter, columns (the header) and rows. JSON: encoding, shape ("array", "object" or
           "lines"), columns (the keys of the object or of the first item) and rows (items). Missing values are null.
"""

import csv
import html
import io
import json
import re
import sys
import unicodedata
import zipfile
from datetime import datetime
from itertools import accumulate, chain, islice

import charset_normalizer
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, NameObject

MB = 1 << 20
MAX_BYTES = 64 * MB  # RUVO caps downloads well below this
MAX_TEXT = 2_000  # characters in one value
MAX_DEPTH = 64  # JSON nesting
MAX_PDF_OBJECTS = 100_000
# Zip-bomb limits, checked against the central directory before any entry is read.
MAX_ENTRIES = 1_000
MAX_UNZIPPED = 100 * MB
MAX_RATIO = 100  # for entries over 1 MB

# Formats refused by their first bytes alone.
REFUSED = {
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1": "a legacy Office file (.xls, .doc or .ppt); only .xlsx workbooks are read",
    b"\x1f\x8b": "a gzip file; compressed files aren't read",
}
ZIP = (b"PK\x03\x04", b"PK\x05\x06")  # an archive, or an empty one
HTML = re.compile(r"<(?:!doctype\s+html|html|head|body)\b", re.I)

# Names that mark active content in a PDF, and the flag each raises. /OpenAction is judged on its own (pdf_flags).
PDF_FLAGS = {
    "/JavaScript": "javascript", "/JS": "javascript",
    "/AA": "auto_action",
    "/Launch": "launch",
    "/EmbeddedFile": "embedded_file", "/EmbeddedFiles": "embedded_file", "/EF": "embedded_file",
    "/RichMedia": "rich_media",
    "/XFA": "xfa",
    "/SubmitForm": "submit_form",
    "/GoToR": "remote_goto",
}
# Each metadata field as pypdf names it in the Info dictionary, and in the XMP metadata.
PDF_META = {
    "title": ("title", "dc_title"),
    "author": ("author", "dc_creator"),
    "subject": ("subject", "dc_description"),
    "keywords": ("keywords", "pdf_keywords"),
    "created": ("creation_date", "xmp_create_date"),
    "modified": ("modification_date", "xmp_modify_date"),
}

WORKBOOK = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
# An absolute path, a drive letter or a ".." step: an entry meant to land outside the folder it is unpacked into.
UNSAFE_PATH = re.compile(r"^[/\\]|^[A-Za-z]:|(?:^|[/\\])\.\.(?:[/\\]|$)")
FORMULA = re.compile(rb"<(?:\w+:)?f\b[^>]*(?<!/)>(.*?)</(?:\w+:)?f>", re.S)
QUOTED = re.compile(r'"[^"]*"')
DDE = re.compile(r"\w\|")  # Excel's DDE syntax: Application|Topic!Item

# A JSON string, or a bracket (captured): nesting is counted on the brackets outside strings.
BRACKETS = re.compile(r'"[^"\\]*(?:\\.[^"\\]*)*"|([\[\]{}])', re.S)
DEPTH = {"[": 1, "{": 1, "]": -1, "}": -1, "": 0}
SPACE = re.compile(r"\s*")
NUMBER = re.compile(r"[-+]?\d[\d.,]*%?")
# Characters that render as nothing or reorder text: zero-width, bidi controls, soft hyphens and C0 controls.
INVISIBLE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u00ad\u200b-\u200f\u202a-\u202e\u2060\u2066-\u2069\ufeff\ud800-\udfff]")


class Rejected(Exception):
    """A document RUVO drops; the message says why."""


def read(data: bytes) -> dict:
    """A document's scan (see the result above)."""
    doc = {"verdict": "clean", "reason": None, "type": None, "flags": [], "meta": {}}
    try:
        scan(data, doc)
    except Rejected as e:
        return rejected(str(e), doc["type"])
    except MemoryError:
        return rejected("too large to read within the sandbox's memory", doc["type"])
    except Exception as e:  # a damaged or hostile file can make a parser fail in any way
        return rejected(f"could not be read ({type(e).__name__}: {clean(str(e))[:200]})", doc["type"])
    if doc["flags"]:
        doc["verdict"] = "risky"
    return doc


def rejected(reason: str, type_: str | None) -> dict:
    return {"verdict": "rejected", "reason": reason, "type": type_, "flags": [], "meta": {}}


def scan(data: bytes, doc: dict) -> None:
    """Finds what the bytes are (their content decides, never a name or declared type), then scans them as that,
    filling in the type, flags and metadata."""
    if not data:
        raise Rejected("an empty file")
    if len(data) > MAX_BYTES:
        raise Rejected(f"larger than {MAX_BYTES // MB} MB")
    if data.startswith(b"%PDF-"):
        doc["type"] = "pdf"
        return scan_pdf(data, doc)
    if data.startswith(ZIP):
        return scan_xlsx(data, doc)
    for magic, what in REFUSED.items():
        if data.startswith(magic):
            raise Rejected(what)
    text, encoding = decode(data)
    head = text[:4096].lstrip()
    if head.startswith("<"):
        if HTML.search(head):
            doc["type"] = "html"
            raise Rejected("a web page, not a document")
        raise Rejected("XML or other markup, not a document")
    if head.startswith(("{", "[")):
        doc["type"] = "json"
        return scan_json(text, encoding, doc)
    return scan_csv(text, encoding, doc)


def scan_pdf(data: bytes, doc: dict) -> None:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        # Its objects can't be read without the password, so there is no metadata to give either.
        doc["flags"] = ["encrypted"]
        return
    doc["flags"] = sorted(pdf_flags(reader))
    doc["meta"] = pdf_meta(reader)


def pdf_flags(reader: PdfReader) -> set[str]:
    """Active content in the objects a viewer can reach from the trailer, those inside object streams included."""
    flags, seen, todo = set(), set(), [reader.trailer]
    while todo:
        obj = todo.pop()
        if isinstance(obj, IndirectObject):
            if (obj.idnum, obj.generation) in seen:
                continue
            seen.add((obj.idnum, obj.generation))
            if len(seen) > MAX_PDF_OBJECTS:
                raise Rejected(f"a PDF with over {MAX_PDF_OBJECTS:,} objects")
            obj = obj.get_object()
        if isinstance(obj, DictionaryObject):  # streams too, without their data
            todo += [*obj.keys(), *obj.values()]
        elif isinstance(obj, ArrayObject):
            todo += obj
        elif isinstance(obj, NameObject) and obj in PDF_FLAGS:
            flags.add(PDF_FLAGS[obj])
    # A destination or a plain /GoTo only sets the first view (LaTeX's hyperref adds one); any other action runs.
    action = reader.trailer["/Root"].get("/OpenAction")
    action = action and action.get_object()
    if isinstance(action, DictionaryObject) and action.get("/S") != "/GoTo":
        flags.add("open_action")
    return flags


def pdf_meta(reader: PdfReader) -> dict:
    """The Info dictionary's properties, with the XMP metadata filling in what it lacks."""
    info, xmp = reader.metadata, attempt(lambda: reader.xmp_metadata)
    meta = {}
    for key, (in_info, in_xmp) in PDF_META.items():
        value = attempt(lambda: getattr(info, in_info)) or attempt(lambda: getattr(xmp, in_xmp))
        if isinstance(value, dict):  # XMP text in several languages
            value = next(iter(value.values()), None)
        elif isinstance(value, list):  # XMP creators
            value = ", ".join(value)
        meta[key] = meta_value(value)
    return {**meta, "pages": len(reader.pages)}


def scan_xlsx(data: bytes, doc: dict) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise Rejected("a damaged ZIP archive")
    check_archive(archive)
    try:
        types = {part.get("ContentType", "") for part in fromstring(archive.read("[Content_Types].xml"))}
    except KeyError:
        types = set()
    if any("macroEnabled" in t for t in types):
        raise Rejected("a macro-enabled workbook (.xlsm)")
    if WORKBOOK not in types:
        raise Rejected("a ZIP archive that isn't an Excel workbook; Word, PowerPoint and compressed files aren't read")
    doc["type"] = "xlsx"
    doc["flags"] = sorted(xlsx_flags(archive))
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    props = workbook.properties
    doc["meta"] = {
        "title": meta_value(props.title),
        "author": meta_value(props.creator),
        "subject": meta_value(props.subject),
        "keywords": meta_value(props.keywords),
        "created": meta_value(props.created),
        "modified": meta_value(props.modified),
        "sheets": [
            {"name": meta_value(ws.title), "hidden": ws.sheet_state != "visible", "rows": ws.max_row, "columns": ws.max_column}
            for ws in workbook.worksheets
        ],
    }


def check_archive(archive: zipfile.ZipFile) -> None:
    """Zip-bomb and path checks on the central directory, before any entry is read. Python's zipfile never unpacks
    more than an entry's declared size, so the declared sizes can be trusted."""
    entries = archive.infolist()
    if len(entries) > MAX_ENTRIES:
        raise Rejected(f"an archive with over {MAX_ENTRIES:,} entries")
    if sum(entry.file_size for entry in entries) > MAX_UNZIPPED:
        raise Rejected(f"a zip bomb: it unpacks to over {MAX_UNZIPPED // MB} MB")
    for entry in entries:
        if entry.file_size > MB and entry.file_size > MAX_RATIO * entry.compress_size:
            raise Rejected(f"a zip bomb: an entry unpacks to over {MAX_RATIO} times its size")
        if UNSAFE_PATH.search(entry.filename):
            raise Rejected(f"an archive entry with an unsafe path ({clean(entry.filename)[:100]})")


def xlsx_flags(archive: zipfile.ZipFile) -> set[str]:
    flags = set()
    for name in archive.namelist():
        lower = name.lower()
        if lower.endswith("vbaproject.bin") or lower.startswith("xl/macrosheets/"):
            flags.add("macros")
        elif lower.startswith(("xl/embeddings/", "xl/activex/")):
            flags.add("ole_object")
        elif lower.startswith("xl/externallinks/") and lower.endswith(".xml"):
            flags.add("dde" if b"ddeLink" in archive.read(name) else "external_link")
        elif lower.endswith(".rels"):
            # A cell's hyperlink is harmless here; any other external target is something Excel would fetch or open.
            for rel in fromstring(archive.read(name)):
                if rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink"):
                    flags.add("external_link")
        elif lower.startswith("xl/worksheets/") and has_dde_formula(archive, name):
            flags.add("dde")
    return flags


def has_dde_formula(archive: zipfile.ZipFile, name: str) -> bool:
    """Whether a worksheet has a DDE formula. Its XML is read a megabyte at a time, so a large sheet isn't held whole;
    each read starts with the end of the one before, so a formula split between them is still found."""
    rest = b""
    with archive.open(name) as part:
        while chunk := part.read(MB):
            xml = rest + chunk
            for formula in FORMULA.findall(xml):
                if DDE.search(QUOTED.sub("", html.unescape(formula.decode("utf-8", "replace")))):
                    return True
            rest = xml[-64 * 1024 :]  # longer than any formula Excel opens
    return False


def decode(data: bytes) -> tuple[str, str]:
    """The document as text and its encoding: UTF-8 when it decodes as UTF-8, else charset-normalizer's best guess."""
    try:
        text, encoding = data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        best = charset_normalizer.from_bytes(data).best()
        if best is None:
            raise Rejected("binary data that isn't a PDF, Excel workbook, CSV or JSON file")
        text, encoding = str(best), best.encoding
    if "\x00" in text:
        raise Rejected("binary data that isn't a PDF, Excel workbook, CSV or JSON file")
    return text.removeprefix("\ufeff"), encoding


def scan_csv(text: str, encoding: str, doc: dict) -> None:
    # A small sample keeps the sniffer's regular expressions quick on any input.
    try:
        dialect = csv.Sniffer().sniff(text[: 16 * 1024], delimiters=",;\t|")
    except csv.Error:
        raise Rejected("text that isn't a CSV or JSON file")
    doc["type"] = "csv"
    header, rows = split_header(enumerate(csv.reader(io.StringIO(text, newline=""), dialect), start=1))
    doc["meta"] = {"encoding": encoding, "delimiter": dialect.delimiter, "columns": names(header), "rows": sum(1 for _ in rows)}


def split_header(rows):
    """A table's header and an iterator over the rows after it, from (row number, values) pairs; empty rows are left
    out. The header is the first of the leading rows that is mostly text and spans at least half the widest of them,
    so a title above it isn't taken for it."""
    rows = ((number, values) for number, values in rows if filled(values))
    first = list(islice(rows, 10))
    if not first:
        return [], iter(())
    width = max(len(filled(values)) for _, values in first)
    for i, (_, values) in enumerate(first):
        cells = filled(values)
        words = [v for v in cells if isinstance(v, str) and not NUMBER.fullmatch(v.strip())]
        if 2 * len(cells) >= width and 2 * len(words) > len(cells):
            break
    else:
        i = 0
    return first[i][1], chain(first[i + 1 :], rows)


def filled(values: list) -> list:
    return [v for v in values if v not in (None, "")]


def scan_json(text: str, encoding: str, doc: dict) -> None:
    if max(accumulate(map(DEPTH.__getitem__, BRACKETS.findall(text))), default=0) > MAX_DEPTH:
        raise Rejected(f"JSON nested over {MAX_DEPTH} levels deep")
    decoder = json.JSONDecoder(parse_constant=not_json)
    start = SPACE.match(text).end()
    if text.startswith("[", start):
        shape, items = "array", array_items(decoder, text, start)
    else:
        value, end = decoder.raw_decode(text, start)
        if SPACE.match(text, end).end() == len(text):
            doc["meta"] = {"encoding": encoding, "shape": "object", "columns": names(value), "rows": None}
            return
        shape, items = "lines", (decoder.decode(line) for line in io.StringIO(text) if line.strip())
    count, columns = 0, None
    for item in items:  # one at a time, so a large file is never held decoded whole
        count += 1
        if columns is None and isinstance(item, dict):
            columns = names(item)
    doc["meta"] = {"encoding": encoding, "shape": shape, "columns": columns or [], "rows": count}


def array_items(decoder: json.JSONDecoder, text: str, start: int):
    """The items of the JSON array at `start`, decoded one by one."""
    i = SPACE.match(text, start + 1).end()
    if not text.startswith("]", i):
        while True:
            item, i = decoder.raw_decode(text, i)
            yield item
            i = SPACE.match(text, i).end()
            if not text.startswith(",", i):
                break
            i = SPACE.match(text, i + 1).end()
    if not text.startswith("]", i):
        raise ValueError(f"expected ',' or ']' at character {i}")
    if SPACE.match(text, i + 1).end() != len(text):
        raise ValueError(f"extra data after character {i}")


def not_json(constant: str):
    raise ValueError(f"{constant} isn't JSON")


def clean(text: str) -> str:
    """NFKC-normalized (ligatures and full-width forms become plain letters), with invisible characters removed."""
    return INVISIBLE.sub("", unicodedata.normalize("NFKC", text))


def meta_value(value) -> str | None:
    """A metadata value as text: dates in ISO 8601, anything else cleaned and capped; None when it is missing."""
    if isinstance(value, datetime):
        return value.isoformat()
    return clean(str(value))[:MAX_TEXT] if value not in (None, "") else None


def names(values) -> list[str]:
    """Column names: a header's values or an object's keys."""
    return [meta_value(value) or "" for value in values]


def attempt(get):
    """A value that may fail to read in a damaged file, as None when it does."""
    try:
        return get()
    except Exception:
        return None


def main() -> None:
    if sys.argv[1:]:
        sys.exit("usage: docparse < document")
    result = read(sys.stdin.buffer.read(MAX_BYTES + 1))
    # A lone surrogate can't be UTF-8, so it goes out as "?".
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8", "replace"))


if __name__ == "__main__":
    main()
